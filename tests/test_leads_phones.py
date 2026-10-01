"""Find phones — the number each person's COMPANY publishes (26-Sep-2026).

Exa's agent reads a company's own site and directory listings and returns the number
it PUBLISHES. That is all it is, and these tests hold the app to saying no more:

  · what a number is read as (a strict format reading — nothing has rung it, and where
    the digits cannot say mobile or landline, "unknown" is the answer, not a guess);
  · one lookup per COMPANY, 15 to a call, only for people with no phone yet, and only
    company details ever sent;
  · the pooled path (the licence server asks Exa, credits are charged per number found)
    and the developer's own-key path, and that a pooled lookup never touches `requests`;
  · what the screen does: Find phones asks in credits first, hands the worker exactly the
    people who need a number, and shows, saves and exports what came back — labelled
    "published, not verified", with the page it came from.

No network, no thread, no key, no real ~/.prism."""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from unittest import mock

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                                    # noqa: E402
from PySide6.QtCore import Qt                                      # noqa: E402
from PySide6.QtWidgets import QApplication, QLabel, QPushButton   # noqa: E402

import core_bridge  # noqa: F401,E402  (puts prism_terminal/core on sys.path)
from addons.leads import person as PP                              # noqa: E402
from prospector import exports as EX                               # noqa: E402
from prospector import gateway as G                                # noqa: E402
from prospector import phones as PH                                # noqa: E402
from prospector.models import Dossier, Lead                        # noqa: E402

from test_leads_cockpit import (_FakeSignal, _Workbench, _dos)     # noqa: E402
from test_leads_credit_pool import FakeServer, _Pooled, _empty_pool   # noqa: E402

_app = QApplication.instance() or QApplication(sys.argv)

MOBILE = "+919825012345"


def lead(name="Asha Rao", company="Acme Works", **kw):
    extra = kw.pop("extra", {"location": "Vadodara"})
    l = Lead(name=name, title="Owner", company=company, **kw)
    l.extra = dict(extra)
    return l


def row(name="Acme Works", phone=MOBILE, kind="mobile", source="own_website",
        url="https://acme.example/contact", domain=""):
    return {"name": name, "domain": domain, "phone": phone, "kind": kind,
            "source_kind": source, "source_url": url, "found_at": "2026-09-26",
            "evidence": "Contact page"}


class TheNumberIsReadStrictly(unittest.TestCase):
    """prospector/phones.normal_phone — the same rules as the licence server's."""

    def test_an_indian_mobile_in_every_way_a_site_writes_it(self):
        for raw in ("+91 98250 12345", "+91-9825012345", "9825012345", "0091 98250 12345",
                    "91 98250 12345", "(+91) 98250-12345"):
            self.assertEqual(PH.normal_phone(raw), (MOBILE, "mobile"), raw)

    def test_a_number_starting_9_is_a_mobile_however_it_is_dialled_or_labelled(self):
        """No Indian STD code starts with 9, so 9… is a mobile — with a trunk 0 too."""
        self.assertEqual(PH.normal_phone("098250 12345"), (MOBILE, "mobile"))
        self.assertEqual(PH.normal_phone("+91 98250 12345", "landline"), (MOBILE, "mobile"))

    def test_a_number_starting_1_to_5_is_a_landline_whatever_the_hint(self):
        for raw in ("+91 265 2345678", "0265-2345678", "022 2345 6789", "011-23456789"):
            self.assertEqual(PH.normal_phone(raw, "mobile")[1], "landline", raw)

    def test_the_writing_of_the_std_code_says_landline_where_the_digits_cannot_decide(self):
        """27-Sep-2026: Siemens' Bengaluru switchboard, printed "+91 (80) 33422000" on its own
        Tel. line, was shown as "Looks like a mobile" — the reader threw the brackets away and
        called every 6-8 number with no label a mobile."""
        for raw in ("+91 (80) 33422000", "+91 80 3342 2000", "+91-80-3342-2000", "080-33422000",
                    "(080) 3342 2000", "+91 (0)80 3342 2000", "0091 80 3342 2000", "80 3342 2000",
                    "+91 79 2658 1234", "0755-2345678", "+91 755 2345678", "(0755) 2345678"):
            self.assertEqual(PH.normal_phone(raw)[1], "landline", raw)
        self.assertEqual(PH.normal_phone("+91 (80) 33422000"), ("+918033422000", "landline"))
        self.assertEqual(PH.normal_phone("+91 (80) 33422000", "mobile")[1], "landline")   # the writing outranks the label

    def test_a_6_to_8_number_with_no_writing_and_no_label_is_unknown_never_a_mobile(self):
        for raw in ("+918033422000", "8033422000", "+91 80334 22000", "07552345678", "+91 78901 23456"):
            self.assertEqual(PH.normal_phone(raw)[1], "unknown", raw)
        self.assertEqual(PH.normal_phone("+91 78901 23456", "mobile"), ("+917890123456", "mobile"))
        self.assertEqual(PH.normal_phone("8033422000", "landline")[1], "landline")

    def test_a_number_saved_before_this_reading_is_not_shown_as_a_mobile_on_a_guess(self):
        """The owner's saved contact: kind "mobile" from the old fallback, no `read` marker."""
        old = lead(phone="+918033422000", extra={"phones": [
            {"number": "+918033422000", "kind": "mobile", "scope": "company", "source": "own_website",
             "source_url": "https://siemens.example/locations", "found_at": "2026-09-27",
             "provider": "exa", "checked": "format"}]})
        self.assertEqual(PH.entry_for(old)["kind"], "unknown")
        self.assertEqual(old.extra["phones"][0]["kind"], "mobile")                # what was saved is not rewritten
        sure = lead(phone=MOBILE, extra={"phones": [{"number": MOBILE, "kind": "mobile", "source": "own_website"}]})
        self.assertEqual(PH.entry_for(sure)["kind"], "mobile")                    # a 9… number was never a guess
        fresh = lead()
        PH.apply(fresh, row(phone="+917890123456", kind="mobile"))                # a page said so: kept
        self.assertEqual(PH.entry_for(fresh)["kind"], "mobile")

    def test_toll_free_and_other_countries(self):
        self.assertEqual(PH.normal_phone("1800 233 4455"), ("+9118002334455", "toll_free"))
        self.assertEqual(PH.normal_phone("+44 20 7946 0958"), ("+442079460958", "unknown"))

    def test_what_is_not_a_number_is_dropped(self):
        for raw in (None, "", "abc", "12345", "2345678", "0000000000", "9999999999",
                    "1234567890", "+91 12345", "+91 98250 12345 ext 22", "x" * 50):
            self.assertIsNone(PH.normal_phone(raw), raw)

    def test_a_mobile_is_shown_the_way_it_is_read_aloud_and_nothing_else_is_split(self):
        self.assertEqual(PH.display(MOBILE, "mobile"), "+91 98250 12345")
        # Where a landline's STD code ends depends on the city; a wrong split is worse than none.
        self.assertEqual(PH.display("+912652345678", "landline"), "+91 2652345678")
        self.assertEqual(PH.display(MOBILE), "+91 9825012345")               # kind not known
        self.assertEqual(PH.display("+442079460958"), "+442079460958")     # as stored

    def test_the_question_put_to_exa_is_the_servers_and_never_asks_it_to_guess(self):
        self.assertIn("Never guess", PH.QUERY)
        self.assertIn("Echo each row's ref back", PH.QUERY)
        body = PH.build_body([{"name": "Acme Works", "domain": "acme.example", "location": "Vadodara"}])
        self.assertEqual(sorted(body), ["effort", "input", "outputSchema", "query"])
        self.assertEqual(body["input"]["data"], [{"ref": "c0", "company": "Acme Works",
                                                  "website": "acme.example",
                                                  "city_or_state": "Vadodara"}])


class TheLeadsAreGroupedByCompany(unittest.TestCase):

    def test_everyone_at_a_company_shares_one_lookup(self):
        a1, a2 = lead("A One"), lead("A Two", company="ACME  works")     # the same firm, written twice
        b = lead("B", company="Beta Tools", extra={"company_domain": "beta.example"})
        [ga, gb] = PH.groups([a1, a2, b])
        self.assertEqual(ga["leads"], [a1, a2])
        self.assertEqual(ga["company"], {"name": "Acme Works", "domain": "", "location": "Vadodara"})
        self.assertEqual(gb["company"]["domain"], "beta.example")
        self.assertEqual(PH.companies_of([a1, a2, b]), 2)

    def test_nobody_with_a_phone_or_no_company_is_looked_up(self):
        has = lead("Has", phone="+91 98250 12345")
        found = lead("Found", extra={"phones": [{"number": "+912652345678"}]})
        nowhere = lead("Nowhere", company="")
        bare = lead("Bare")
        self.assertEqual([g["leads"] for g in PH.groups([has, found, nowhere, bare])], [[bare]])
        self.assertFalse(PH.needs_phone(has) or PH.needs_phone(found) or PH.needs_phone(nowhere))
        self.assertTrue(PH.needs_phone(bare))

    def test_a_company_can_be_known_only_by_its_website(self):
        l = lead("W", company="", extra={"website": "https://www.Beta.example/about"})
        self.assertTrue(PH.needs_phone(l))
        [g] = PH.groups([l])
        self.assertEqual(g["company"], {"name": "beta.example", "domain": "beta.example", "location": ""})

    def test_a_number_is_recorded_with_where_it_came_from_and_never_overwrites_one_they_had(self):
        l = lead()
        self.assertTrue(PH.apply(l, row()))
        self.assertEqual(l.phone, MOBILE)
        found = PH.entry_for(l)
        self.assertEqual((found["kind"], found["scope"], found["source"], found["provider"],
                          found["checked"], found["found_at"], found["source_url"]),
                         ("mobile", "company", "own_website", "exa", "format", "2026-09-26",
                          "https://acme.example/contact"))
        self.assertFalse(PH.apply(l, row()))                         # the same number twice adds nothing
        self.assertEqual(len(l.extra["phones"]), 1)
        theirs = lead("Sheet", phone="+91 99999 88888")
        PH.apply(theirs, row(phone="+912652345678", kind="landline"))
        self.assertEqual(theirs.phone, "+91 99999 88888")            # what they came with stays
        self.assertIsNone(PH.entry_for(theirs))                      # …and nothing is claimed about it
        self.assertEqual(theirs.extra["phones"][0]["number"], "+912652345678")

    def test_an_answer_for_the_wrong_company_is_not_used(self):
        got = PH.read_rows([{"name": "Acme", "domain": "", "location": ""}],
                           [{"ref": "c7", "phone": MOBILE}])
        self.assertEqual(got, [])

    def test_exa_calling_a_printed_landline_a_mobile_does_not_make_it_one(self):
        """The Siemens switchboard, "+91 (80) 33422000", with Exa's label "mobile"."""
        batch = [{"name": "Siemens Ltd", "domain": "", "location": "Bengaluru"},
                 {"name": "Other Co", "domain": "", "location": ""}]
        got = PH.read_rows(batch, [{"ref": "c0", "phone": "+91 (80) 33422000", "phone_type": "mobile",
                                    "source_kind": "own_website", "source_url": "https://siemens.example/x"},
                                   {"ref": "c1", "phone": "+918033422001"}])
        self.assertEqual([(r["phone"], r["kind"]) for r in got],
                         [("+918033422000", "landline"), ("+918033422001", "unknown")])


@pytest.mark.pooled
class FindingThemOnThePool(_Pooled):
    """The licence server asks Exa with its key; the customer's credits pay per number found."""

    CFG = {"exa_api_key": G.POOL_KEY}

    def test_one_lookup_per_company_for_only_the_people_who_need_a_number(self):
        a1, a2, has, nowhere = lead("A1"), lead("A2"), lead("Has", "Beta Tools", phone="+91 98250 11111"), \
            lead("N", company="")
        self.server.answer("phones", {"results": [row()], "asked": 1}, charged=2)
        out = PH.find([a1, a2, has, nowhere], self.CFG)
        [(op, body)] = [c for c in self.server.calls if c[0] == "phones"]
        self.assertEqual(body["companies"], [{"name": "Acme Works", "domain": "", "location": "Vadodara"}])
        self.assertEqual(out, {"asked": 1, "found": 1, "given": 2, "failed": 0, "why": []})
        self.assertEqual((a1.phone, a2.phone), (MOBILE, MOBILE))     # both people at the firm
        self.assertEqual((has.phone, nowhere.phone), ("+91 98250 11111", ""))
        self.assertEqual(G.used(), 2)                                # and it was charged to the pool

    def test_fifteen_companies_to_a_call(self):
        leads = [lead("P%d" % i, "Company %d" % i) for i in range(20)]
        self.server.answer("phones", {"results": [], "asked": 0}, charged=0)
        out = PH.find(leads, self.CFG)
        sizes = [len(b["companies"]) for op, b in self.server.calls if op == "phones"]
        self.assertEqual(sizes, [15, 5])
        self.assertEqual(out, {"asked": 20, "found": 0, "given": 0, "failed": 0, "why": []})

    def test_a_company_that_publishes_nothing_is_asked_but_not_found_and_nobody_changes(self):
        a, b = lead("A"), lead("B", "Beta Tools")
        self.server.answer("phones", {"results": [row("Beta Tools", phone="+912652345678", kind="landline")],
                                      "asked": 2}, charged=2)
        out = PH.find([a, b], self.CFG)
        self.assertEqual(out, {"asked": 2, "found": 1, "given": 1, "failed": 0, "why": []})
        self.assertEqual((a.phone, b.phone), ("", "+912652345678"))
        self.assertFalse(a.extra.get("phones"))

    def test_a_failed_lookup_changes_nothing_and_says_why_in_the_servers_words(self):
        a = lead()
        self.server.fail("phones", G.GatewayError("PROVIDER_ERROR", "Exa was unreachable."))
        out = PH.find([a], self.CFG)
        self.assertEqual(out, {"asked": 0, "found": 0, "given": 0, "failed": 1,
                               "why": ["Exa was unreachable."]})
        self.assertEqual((a.phone, a.extra.get("phones")), ("", None))

    def test_a_server_that_is_too_old_is_said_plainly_not_as_a_failed_lookup(self):
        a = lead()
        self.server.fail("phones", G.GatewayError("SERVER_OUTDATED", G.OUTDATED_MESSAGE))
        out = PH.find([a], self.CFG)
        self.assertEqual((out["asked"], out["failed"], out["why"]), (0, 1, [G.OUTDATED_MESSAGE]))

    def test_the_same_reason_for_every_batch_is_said_once(self):
        leads = [lead("P%d" % i, "Company %d" % i) for i in range(20)]
        self.server.fail("phones", G.GatewayError("PROVIDER_ERROR", "The data provider didn't answer."))
        out = PH.find(leads, self.CFG)
        self.assertEqual((out["failed"], out["why"]), (20, ["The data provider didn't answer."]))

    def test_a_pool_that_runs_dry_stops_the_run_and_keeps_what_was_found(self):
        leads = [lead("P%d" % i, "Company %d" % i) for i in range(20)]
        first = [row("Company %d" % i, phone="+9198250%05d" % (10000 + i)) for i in range(15)]
        calls = []

        def server(path, body):
            op = path.rsplit("/", 1)[-1]
            if op != "phones":
                return {"balance": 3, "rates": {}, "ready": {}}
            calls.append(len(body["companies"]))
            if len(calls) == 1:
                return {"result": {"results": first, "asked": 15}, "charged": 30, "balance": 3}
            raise _empty_pool()
        G.set_transport(server)
        out = PH.find(leads, self.CFG)
        self.assertEqual(calls, [15, 5])                             # the second was refused…
        self.assertEqual((out["found"], out["given"], out["failed"]), (15, 15, 5))
        self.assertEqual(sum(1 for l in leads if l.phone), 15)       # …and the first's numbers were kept
        calls.clear()
        PH.find([lead("Again", "Another Co")], self.CFG)              # the pool is known to be empty:
        self.assertEqual(calls, [])                                   # no further call is even made

    def test_only_company_details_are_sent(self):
        l = lead("Secret Person", extra={"location": "Vadodara", "linkedin": "https://linkedin.com/in/x",
                                         "company_domain": "acme.example"})
        l.email = "secret.person@acme.example"
        self.server.answer("phones", {"results": [], "asked": 1}, charged=0)
        PH.find([l], self.CFG)
        [(op, body)] = [c for c in self.server.calls if c[0] == "phones"]
        sent = repr(body)
        for private in ("Secret Person", "secret.person", "linkedin"):
            self.assertNotIn(private, sent)
        self.assertEqual(body["companies"][0]["domain"], "acme.example")


class FindingThemWithYourOwnKey(unittest.TestCase):
    """The developer's PRISM_LEADS_DIRECT=1: the same lookup, asked of Exa directly."""

    CFG = {"exa_api_key": "live-exa-key"}

    def run_of(self, answers):
        return {"status": "completed", "output": {"structured": {"companies": answers}}}

    def test_it_asks_exa_and_never_the_gateway(self):
        a = lead()
        answers = [{"ref": "c0", "phone": "0265 2345678", "phone_type": "landline",
                    "source_kind": "directory", "source_url": "https://dir.example/acme"}]
        with mock.patch.object(PH, "_agent_run", return_value=self.run_of(answers)) as run, \
                mock.patch.object(G, "company_phones", side_effect=AssertionError("asked the pool")):
            out = PH.find([a], self.CFG)
        self.assertEqual(run.call_args.args[1], "live-exa-key")
        self.assertEqual(out, {"asked": 1, "found": 1, "given": 1, "failed": 0, "why": []})
        self.assertEqual((a.phone, PH.entry_for(a)["kind"], PH.entry_for(a)["source"]),
                         ("+912652345678", "landline", "directory"))

    def test_a_run_that_could_not_be_had_counts_as_failed(self):
        a = lead()
        with mock.patch.object(PH, "_agent_run", return_value=None):
            out = PH.find([a], self.CFG)
        self.assertEqual((out["failed"], out["found"], a.phone), (1, 0, ""))

    def test_the_agent_run_is_started_then_waited_for(self):
        sent = []

        class Reply:
            def __init__(self, status, body):
                self.status_code, self._body = status, body

            def json(self):
                return self._body
        polls = iter([{"status": "running"}, {"status": "completed", "output": {"x": 1}}])

        def post(url, headers=None, json=None, timeout=None):
            sent.append((url, headers, json))
            return Reply(200, {"id": "run_abc123"})

        def get(url, headers=None, timeout=None):
            sent.append((url, headers, None))
            return Reply(200, next(polls))
        import requests
        with mock.patch.object(requests, "post", post), mock.patch.object(requests, "get", get), \
                mock.patch.object(PH, "_sleep", lambda s: None):
            run = PH._agent_run({"query": "q"}, "live-exa-key")
        self.assertEqual(run["status"], "completed")
        self.assertEqual([u for u, _h, _b in sent],
                         [PH.EXA_AGENT_URL, PH.EXA_AGENT_URL + "/run_abc123", PH.EXA_AGENT_URL + "/run_abc123"])
        self.assertEqual(sent[0][1]["x-api-key"], "live-exa-key")

    def test_a_run_that_fails_or_never_ends_or_has_a_hostile_id_is_none(self):
        class Reply:
            def __init__(self, status, body):
                self.status_code, self._body = status, body

            def json(self):
                return self._body
        import requests
        cases = (("failed", {"id": "run_abc123"}, {"status": "failed"}),
                 ("stalled", {"id": "run_abc123"}, {"status": "running"}),
                 ("hostile id", {"id": "../../admin"}, {"status": "completed"}),
                 ("no id", {}, {"status": "completed"}))
        for label, started, polled in cases:
            with mock.patch.object(requests, "post", lambda *a, **k: Reply(200, started)), \
                    mock.patch.object(requests, "get", lambda *a, **k: Reply(200, polled)), \
                    mock.patch.object(PH, "_sleep", lambda s: None):
                self.assertIsNone(PH._agent_run({}, "k", max_wait=6.0, poll=3.0), label)


class _Reply:
    """A `requests` response: a status and a JSON body (None = not JSON)."""

    def __init__(self, status, body=None):
        self.status_code, self._body = status, body

    def json(self):
        if self._body is None:
            raise ValueError("not json")
        return self._body


EXA_NO_CREDIT = {"requestId": "c0b29547aab81bc8c254d950119b4af2", "tag": "NO_MORE_CREDITS",
                 "error": "You have exceeded your credits limit. Please top up to keep using Exa "
                          "at dashboard.exa.ai"}                      # what Exa really sent (26-Sep-2026)


class WhenExaSaysNoTheOwnerIsToldWhy(unittest.TestCase):
    """26-Sep-2026: Find their phone… said "Found 0 phone number(s) for 0 companies · 1 couldn't be
    looked up" and nothing more, while Exa had answered 402 — this account is out of credit.
    Every way a lookup fails now says which, in a sentence a person can act on."""

    KEY = "live-exa-key"

    def why(self, post, get=None, **kw):
        import requests
        why = []
        running = lambda *a, **k: _Reply(200, {"status": "running"})       # noqa: E731
        with mock.patch.object(requests, "post", post), \
                mock.patch.object(requests, "get", get or running), \
                mock.patch.object(PH, "_sleep", lambda s: None):
            self.assertIsNone(PH._agent_run({}, self.KEY, why=why, **kw))
        return why

    def test_an_exa_account_with_no_credit_says_so_and_where_to_top_up(self):
        [said] = self.why(lambda *a, **k: _Reply(402, EXA_NO_CREDIT))
        self.assertEqual(said, "Exa says this account is out of credits — top it up at dashboard.exa.ai.")

    def test_a_key_exa_rejects_points_at_where_the_key_is_kept(self):
        for status in (401, 403):
            [said] = self.why(lambda *a, **k: _Reply(status, {"error": "invalid api key"}))
            self.assertIn("did not accept the API key (HTTP %d)" % status, said)
            self.assertIn("Search settings", said)

    def test_a_rate_limit_and_an_exa_outage_are_told_apart_from_a_refusal(self):
        [limited] = self.why(lambda *a, **k: _Reply(429, {"error": "slow down"}))
        self.assertIn("limiting this key", limited)
        [down] = self.why(lambda *a, **k: _Reply(503, None))
        self.assertIn("problem on its side (HTTP 503)", down)

    def test_a_request_exa_does_not_accept_quotes_exa_and_never_more_than_a_line_of_it(self):
        [said] = self.why(lambda *a, **k: _Reply(400, {"error": "outputSchema is invalid " + "x" * 400}))
        self.assertTrue(said.startswith("Exa did not start the lookup (HTTP 400): outputSchema is invalid"))
        self.assertLess(len(said), 220)
        [bare] = self.why(lambda *a, **k: _Reply(400, None))              # a body that is not JSON
        self.assertEqual(bare, "Exa did not start the lookup (HTTP 400).")

    def test_a_reply_with_no_run_to_follow_is_said(self):
        [said] = self.why(lambda *a, **k: _Reply(200, {}))
        self.assertIn("no run Prism can follow", said)

    def test_a_run_that_goes_wrong_while_it_is_waited_for(self):
        started = lambda *a, **k: _Reply(200, {"id": "run_abc123"})       # noqa: E731
        [gone] = self.why(started, lambda *a, **k: _Reply(404, {"error": "no such run"}))
        self.assertEqual(gone, "Exa did not return the lookup's answer (HTTP 404): no such run")
        [spent] = self.why(started, lambda *a, **k: _Reply(402, EXA_NO_CREDIT))
        self.assertIn("out of credits", spent)
        [failed] = self.why(started, lambda *a, **k: _Reply(200, {"status": "failed", "error": "schema rejected"}))
        self.assertEqual(failed, "Exa's lookup failed: schema rejected")
        [cancelled] = self.why(started, lambda *a, **k: _Reply(200, {"status": "cancelled"}))
        self.assertEqual(cancelled, "Exa's lookup cancelled.")
        [slow] = self.why(started, max_wait=6.0, poll=3.0)               # the default poll answers "running"
        self.assertEqual(slow, "Exa did not finish within 6 seconds.")

    def test_no_connection_says_so_without_the_exception_text(self):
        def post(*a, **k):
            raise ConnectionError("https://api.exa.ai/agent/runs?x-api-key=SECRET-IN-A-URL")
        [said] = self.why(post)
        self.assertEqual(said, "Couldn't reach Exa (ConnectionError) — check the internet connection.")
        self.assertNotIn("SECRET", said)

    def test_the_lookup_the_owner_pressed_reports_it_instead_of_found_zero_of_zero(self):
        a = lead("Pankaj Vyas", "Siemens")
        import requests
        with mock.patch.object(requests, "post", lambda *x, **k: _Reply(402, EXA_NO_CREDIT)), \
                mock.patch.object(G, "company_phones", side_effect=AssertionError("asked the pool")):
            out = PH.find([a], {"exa_api_key": self.KEY})
        self.assertEqual(out, {"asked": 0, "found": 0, "given": 0, "failed": 1, "why": [
            "Exa says this account is out of credits — top it up at dashboard.exa.ai."]})
        self.assertEqual((a.phone, a.extra.get("phones")), ("", None))
        self.assertNotIn(self.KEY, repr(out))

    def test_twenty_companies_that_all_hit_the_same_wall_give_one_reason(self):
        leads = [lead("P%d" % i, "Company %d" % i) for i in range(20)]
        import requests
        with mock.patch.object(requests, "post", lambda *x, **k: _Reply(402, EXA_NO_CREDIT)):
            out = PH.find(leads, {"exa_api_key": self.KEY})
        self.assertEqual((out["asked"], out["failed"], len(out["why"])), (0, 20, 1))

    def test_an_answer_that_finished_but_cannot_be_read_is_said(self):
        a = lead()
        with mock.patch.object(PH, "_agent_run", return_value={"status": "completed",
                                                              "output": {"structured": None}}):
            out = PH.find([a], {"exa_api_key": self.KEY})
        self.assertEqual((out["asked"], out["failed"]), (0, 1))
        self.assertIn("not a list of companies", out["why"][0])


class TheWorkerRunsTheLookup(unittest.TestCase):
    def test_it_hands_the_leads_to_the_engine_and_reports_the_counts(self):
        from addons.leads.workers import LeadsPhoneWorker
        leads = [lead()]
        counts, progress, failed = [], [], []
        w = LeadsPhoneWorker(leads, {"exa_api_key": "live"})
        w.done.connect(counts.append)
        w.progress.connect(lambda i, n: progress.append((i, n)))
        w.failed.connect(failed.append)
        with mock.patch.object(PH, "find", side_effect=lambda ls, cfg, on_progress=None: (
                on_progress(1, 1), {"asked": 1, "found": 1, "given": 1, "failed": 0})[1]) as find:
            w.run()
        self.assertEqual((counts, progress, failed), ([{"asked": 1, "found": 1, "given": 1, "failed": 0}],
                                                       [(1, 1)], []))
        self.assertEqual(find.call_args.args[0], leads)

    def test_a_crash_is_a_failure_message_not_an_exception(self):
        from addons.leads.workers import LeadsPhoneWorker
        failed = []
        w = LeadsPhoneWorker([lead()], {})
        w.failed.connect(failed.append)
        with mock.patch.object(PH, "find", side_effect=RuntimeError("boom")):
            w.run()
        self.assertEqual(failed, ["boom"])


class _FakePhoneWorker:
    """Stands in for LeadsPhoneWorker: records what it was handed; start() only records."""
    made: list = []

    def __init__(self, leads, cfg):
        self.leads, self.cfg, self.started = leads, cfg, False
        self.progress, self.done, self.failed = _FakeSignal(), _FakeSignal(), _FakeSignal()
        _FakePhoneWorker.made.append(self)

    def start(self):
        self.started = True


class FindPhonesOnTheScreen(_Workbench):
    """The bulk bar's Find phones and the person panel's Find their phone…, through the workbench."""

    def _people(self, wb, n=3):
        from prospector.engine import RunResult
        leads = [Lead(name=f"P{i} Singh", title="Plant Head", company=f"Co {i}", fit_score=60 + i)
                 for i in range(n)]
        for l in leads:
            l.extra = {"location": "Vadodara"}
        res = RunResult(dossiers=[], total_in_sheet=n, signal_source="", all_leads=leads)
        wb._start_export = lambda *a, **k: None
        wb._next_mode = "icp_leads_only"
        wb._next_params = {"mode": "icp_leads_only", "offer": "Automation"}
        wb._jobs = 1
        wb._on_prepared(res, [])
        return leads

    def _tick(self, wb, leads):
        ck = wb._cockpit.leads
        for r in range(ck._table.rowCount()):
            if ck._dossier_at(r).lead in leads:
                ck._table.item(r, 0).setCheckState(Qt.Checked)

    def _patched_worker(self):
        _FakePhoneWorker.made = []
        orig = self._WB.LeadsPhoneWorker
        self._WB.LeadsPhoneWorker = _FakePhoneWorker
        self.addCleanup(setattr, self._WB, "LeadsPhoneWorker", orig)

    def _bench(self):
        wb = self._WB.LeadsWorkbench({"exa_api_key": "live-exa"})
        wb._confirm_phones = lambda n: True
        self._patched_worker()
        return wb

    def test_the_worker_gets_exactly_the_people_who_need_a_number(self):
        wb = self._bench()
        leads = self._people(wb)
        leads[1].phone = "+91 98250 12345"                      # this one came with a number
        self._tick(wb, leads)
        wb._find_phones(wb._cockpit.leads.selected())
        [worker] = _FakePhoneWorker.made
        self.assertTrue(worker.started)
        self.assertEqual(sorted(l.name for l in worker.leads), ["P0 Singh", "P2 Singh"])   # not P1
        self.assertIn("2 companies", wb._status.text())

    def test_a_second_press_while_it_runs_starts_nothing(self):
        wb = self._bench()
        leads = self._people(wb)
        self._tick(wb, leads)
        picked = wb._cockpit.leads.selected()
        wb._find_phones(picked)
        wb._find_phones(picked)
        self.assertEqual(len(_FakePhoneWorker.made), 1)

    def test_the_confirm_counts_companies_not_people(self):
        wb = self._WB.LeadsWorkbench({"exa_api_key": "live-exa"})
        asked = []
        wb._confirm_phones = lambda n: (asked.append(n), False)[1]
        self._patched_worker()
        leads = self._people(wb)
        leads[1].company = leads[0].company                     # two people, one company
        self._tick(wb, leads)
        wb._find_phones(wb._cockpit.leads.selected())
        self.assertEqual(asked, [2])
        self.assertEqual(_FakePhoneWorker.made, [])             # declined: nothing ran, nothing spent

    @pytest.mark.pooled
    def test_on_the_pool_it_asks_in_credits_and_says_what_it_is(self):
        wb = self._WB.LeadsWorkbench({"exa_api_key": "live-exa"})
        self._patched_worker()
        G.set_transport(FakeServer(balance=50))
        self.addCleanup(G.set_transport, None)
        self.addCleanup(G.forget)
        with mock.patch.object(G, "_auth_body", lambda **e: {"license_id": "L", "device_fp": "D", **e}):
            wb._on_credits(G.status())
            said = []
            wb._confirm_pool = lambda title, lines: (said.append((title, lines)), False)[1]
            leads = self._people(wb)
            self._tick(wb, leads)
            wb._find_phones(wb._cockpit.leads.selected())
        [(title, lines)] = said
        self.assertEqual(title, "Find phones")
        self.assertIn("3 companies", lines[0])
        self.assertIn("up to 6 credits", lines[0])              # 3 x the 2-credit default
        self.assertIn("only for a number found", lines[1])
        self.assertIn("not a direct dial", lines[1])            # what it is, before it is bought
        self.assertEqual(_FakePhoneWorker.made, [])

    @pytest.mark.pooled
    def test_a_lookup_the_server_has_not_switched_on_says_so_and_charges_nothing(self):
        wb = self._WB.LeadsWorkbench({"exa_api_key": "live-exa"})
        self._patched_worker()
        server = FakeServer(balance=50)
        server.__class__ = type("S", (FakeServer,), {"__call__": lambda self, path, body: (
            {"balance": 50, "rates": {}, "ready": {"phones": False}} if path.endswith("status")
            else FakeServer.__call__(self, path, body))})
        G.set_transport(server)
        self.addCleanup(G.set_transport, None)
        self.addCleanup(G.forget)
        with mock.patch.object(G, "_auth_body", lambda **e: {"license_id": "L", "device_fp": "D", **e}):
            wb._on_credits(G.status())
            leads = self._people(wb)
            self._tick(wb, leads)
            wb._find_phones(wb._cockpit.leads.selected())
        self.assertIn("isn't switched on yet", wb._status.text())
        self.assertEqual(_FakePhoneWorker.made, [])

    def test_with_your_own_keys_it_needs_an_exa_key(self):
        wb = self._WB.LeadsWorkbench({})
        wb._confirm_phones = lambda n: True
        self._patched_worker()
        leads = self._people(wb)
        self._tick(wb, leads)
        wb._find_phones(wb._cockpit.leads.selected())
        self.assertIn("needs an Exa API key", wb._status.text())
        self.assertEqual(_FakePhoneWorker.made, [])

    def test_nobody_left_to_look_up_says_so(self):
        wb = self._bench()
        leads = self._people(wb)
        for l in leads:
            l.phone = "+91 98250 12345"
        self._tick(wb, leads)
        wb._find_phones(wb._cockpit.leads.selected())
        self.assertIn("already have a phone number", wb._status.text())
        self.assertEqual(_FakePhoneWorker.made, [])

    def test_what_comes_back_is_shown_saved_and_labelled_published_not_verified(self):
        from addons.leads import sessions
        wb = self._bench()
        leads = self._people(wb)
        PH.apply(leads[0], row("Co 0"))
        PH.apply(leads[1], row("Co 1", phone="+912652345678", kind="landline", source="directory"))
        wb._jobs = 1
        wb._on_phones_found({"asked": 3, "found": 2, "given": 2, "failed": 0})
        said = wb._status.text()
        self.assertIn("Found 2 phone number(s) for 3 companies", said)
        self.assertIn("published, not verified", said)
        self.assertIn("1 publish none", said)
        self.assertEqual(wb._jobs, 0)
        loaded = sessions.load(self._sessions, wb._session_id)
        self.assertEqual([l.phone for l in loaded["all_leads"]][:2], [MOBILE, "+912652345678"])
        self.assertEqual(loaded["all_leads"][0].extra["phones"][0]["checked"], "format")   # provenance kept

    def test_a_lookup_that_partly_failed_says_how_many(self):
        wb = self._bench()
        self._people(wb)
        wb._jobs = 1
        wb._on_phones_found({"asked": 1, "found": 1, "given": 1, "failed": 4})
        self.assertIn("4 couldn't be looked up", wb._status.text())

    def test_a_lookup_that_never_ran_says_why_and_is_not_found_zero_of_zero(self):
        """The owner's own case (26-Sep-2026): Exa had no credit left, and the line read
        "Found 0 phone number(s) for 0 companies · 1 couldn't be looked up"."""
        wb = self._bench()
        self._people(wb)
        wb._jobs = 1
        wb._on_phones_found({"asked": 0, "found": 0, "given": 0, "failed": 1, "why": [
            "Exa says this account is out of credits — top it up at dashboard.exa.ai."]})
        said = wb._status.text()
        self.assertEqual(said, "No phone numbers were looked up.   ·   Exa says this account is out of "
                               "credits — top it up at dashboard.exa.ai.")
        self.assertNotIn("Found 0", said)
        self.assertEqual(wb._jobs, 0)

    def test_a_lookup_that_never_ran_still_says_so_when_no_reason_is_known(self):
        wb = self._bench()
        self._people(wb)
        wb._jobs = 1
        wb._on_phones_found({"asked": 0, "found": 0, "given": 0, "failed": 2})
        self.assertEqual(wb._status.text(), "No phone numbers were looked up.")

    def test_a_partial_failure_carries_its_reason_after_its_count(self):
        wb = self._bench()
        self._people(wb)
        wb._jobs = 1
        wb._on_phones_found({"asked": 1, "found": 1, "given": 1, "failed": 4, "why": [
            "Exa is limiting this key for the moment (HTTP 429) — try again in a minute."]})
        said = wb._status.text()
        self.assertIn("Found 1 phone number(s) for 1 companies — published, not verified.", said)
        self.assertIn("4 couldn't be looked up. Exa is limiting this key", said)

    @pytest.mark.pooled
    def test_the_pools_out_of_credits_line_is_said_once_not_twice(self):
        wb = self._WB.LeadsWorkbench({"exa_api_key": "live-exa"})
        self._people(wb)
        wb._jobs = 1
        with mock.patch.object(G, "exhausted", return_value="You've run out of credits."):
            wb._on_phones_found({"asked": 0, "found": 0, "given": 0, "failed": 2,
                                 "why": ["You've run out of credits."]})
        said = wb._status.text()
        self.assertNotIn("You've run out of credits", said)      # _used_note says it, in the pool's words
        self.assertIn("stopped: out of credits", said)
        self.assertTrue(said.startswith("No phone numbers were looked up."))

    def test_the_person_panel_asks_for_one_person_through_the_enrichment_queue(self):
        wb = self._bench()
        leads = self._people(wb)
        dos = wb._cockpit.leads._dossier_at(0)
        wb._panel_enrich(dos, ["phone"])
        [worker] = _FakePhoneWorker.made
        self.assertEqual([l.name for l in worker.leads], [dos.lead.name])


class ThePersonPanelSaysWhatTheNumberIs(unittest.TestCase):

    def panel(self, l):
        p = PP.PersonPanel()
        p.resize(760, 900)
        self.asked, self.opened = [], []
        p.enrichRequested.connect(lambda row, picks: self.asked.append(picks))
        p.open_url = self.opened.append                             # the test seam: no browser
        p.show_person(PP.PersonView(row=Dossier(lead=l, status="ok", verdict="warm")))
        return p

    @staticmethod
    def texts(p):
        return [w.text() for w in p._left.findChildren(QLabel)]

    def test_a_found_number_shows_its_kind_its_source_and_that_it_is_not_verified(self):
        l = lead()
        PH.apply(l, row())
        p = self.panel(l)
        texts = self.texts(p)
        self.assertIn("+91 98250 12345", texts)
        self.assertIn("Looks like a mobile", texts)                 # by format, never "mobile" as a fact
        self.assertIn("The company's number, published on their website · found 26 Sep 2026 · "
                      "not verified", texts)
        link = [b for b in p._left.findChildren(QPushButton) if b.text() == "Open the page it came from"]
        self.assertEqual(len(link), 1)
        link[0].click()
        self.assertEqual(self.opened, ["https://acme.example/contact"])
        self.assertFalse([b for b in p._left.findChildren(QPushButton) if b.text() == "Find their phone…"])

    def test_an_office_line_is_not_shown_as_a_mobile_neither_new_nor_saved_before_the_fix(self):
        """27-Sep-2026, the owner's screenshot: Siemens' Bengaluru switchboard, "+91 80334 22000",
        "Looks like a mobile". A new lookup reads it as the landline it is; the contact saved
        under the old rules is shown as a plain phone number, not grouped like a mobile."""
        fresh = lead()
        PH.apply(fresh, row(phone="+918033422000", kind="landline"))
        saved = lead(phone="+918033422000", extra={"phones": [
            {"number": "+918033422000", "kind": "mobile", "scope": "company", "source": "own_website",
             "source_url": "https://siemens.example/locations", "found_at": "2026-09-27",
             "provider": "exa", "checked": "format"}]})
        for l, pill in ((fresh, "Office / landline line"), (saved, "Phone number")):
            texts = self.texts(self.panel(l))
            self.assertIn("+91 8033422000", texts)
            self.assertIn(pill, texts)
            self.assertNotIn("Looks like a mobile", texts)
            self.assertNotIn("+91 80334 22000", texts)                          # the mobile grouping

    def test_a_number_they_came_with_claims_nothing_about_its_source(self):
        p = self.panel(lead(phone="+91 98250 12345"))
        texts = self.texts(p)
        self.assertIn("+91 98250 12345", texts)
        self.assertFalse([t for t in texts if "published" in t or "not verified" in t])

    def test_no_phone_offers_to_find_it_and_asks_for_it(self):
        p = self.panel(lead())
        self.assertIn("No phone yet", self.texts(p))
        [find] = [b for b in p._left.findChildren(QPushButton) if b.text() == "Find their phone…"]
        find.click()
        self.assertEqual(self.asked, [["phone"]])

    def test_a_person_at_no_company_is_not_offered_a_lookup(self):
        p = self.panel(lead(company=""))
        self.assertFalse([b for b in p._left.findChildren(QPushButton) if b.text() == "Find their phone…"])
        self.assertFalse(p.enrich_boxes["phone"].isEnabled())


class TheSheetSaysWhereTheNumberCameFrom(unittest.TestCase):

    def test_the_phone_source_column(self):
        found = lead()
        PH.apply(found, row(source="directory"))
        self.assertEqual(EX.phone_source_of(found),
                         "The company's number — a business directory, 26 Sep 2026, not verified")
        self.assertEqual(EX.phone_source_of(lead(phone="+91 98250 12345")), "Your sheet")
        self.assertEqual(EX.phone_source_of(lead()), "")

    def test_the_leads_sheet_carries_it_beside_the_number(self):
        found = lead()
        PH.apply(found, row())
        with tempfile.TemporaryDirectory() as tmp:
            path = EX.leads_xlsx([found, lead("Other", "Beta Tools")], os.path.join(tmp, "leads.xlsx"))
            ws = openpyxl.load_workbook(path).worksheets[0]
            head = [c.value for c in ws[1]]
            at = head.index("Phone No.")
            self.assertEqual(head[at + 1], "Phone source")
            rows = {r[3].value: r for r in ws.iter_rows(min_row=2)}
            self.assertEqual(rows["Asha Rao"][at].value, MOBILE)
            self.assertIn("not verified", rows["Asha Rao"][at + 1].value)
            self.assertIn(rows["Other"][at].value, (None, ""))


if __name__ == "__main__":
    unittest.main()
