"""Quote by product code: the rate list from any file, every code in the
mail with its quantity, and the opt-in automatic send.

The client's ask (11 Sep 2026): customers write the catalogue code and the
pieces; the code is on the price list; quote them straight away. Pinned:

  · quoting.find_requests reads every rate-list code a mail names, in
    order, with the quantity written beside it in the ways people write
    it -- and returns a code with no quantity as NOT confident;
  · a rate list loads from a Word file and from plain text, not only CSV;
  · the setup screen saves the switch, off by default;
  · the quotation window shows a lines table for two codes and prefills
    the form for one;
  · with the switch on, a check that brings in a confident inquiry sends
    the quotation and marks the row Quoted; a mail with a code but no
    quantity is held with the reason in Notes; with the switch off nothing
    is sent.

No SMTP, no IMAP, no ~/.prism: SendWorker is a fake that answers at once,
config.save is refused as in test_inquiry_ui, and every folder is a temp.
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from decimal import Decimal
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("PRISM_LICENSE_OFFLINE_DEV", "1")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from PySide6.QtCore import QObject, Signal  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

import core_bridge as CB  # noqa: E402
from core import mailflow, register  # noqa: E402
from addons.inquiry import autoquote  # noqa: E402
from addons.inquiry import dialog as UI  # noqa: E402
from addons.inquiry.setup import InquirySetupDialog  # noqa: E402
from test_inquiry_ui import _NoSave, message, ready_cfg  # noqa: E402

_app = QApplication.instance() or QApplication([])
quoting = CB.get_quoting()

RATES = ("Code,Description,Unit,Rate\n"
         "1128K,Office chair mesh back,nos,250\n"
         "1129K,Office chair leather,nos,400\n"
         "T-55,Meeting table 6 seat,nos,9000\n")


def _rates(folder: str) -> str:
    p = os.path.join(folder, "rates.csv")
    with open(p, "w", encoding="utf-8") as f:
        f.write(RATES)
    return p


class TheCodesInAMail(unittest.TestCase):

    def setUp(self):
        self.items = quoting.load_rates(_rates(tempfile.mkdtemp()))

    def _codes(self, text):
        return [(r.item.code, r.quantity, r.confident)
                for r in quoting.find_requests(text, self.items)]

    def test_every_way_people_write_a_quantity(self):
        self.assertEqual(self._codes("need 40 pcs of chair 1128k and 1129K x 10"),
                         [("1128K", Decimal("40"), True), ("1129K", Decimal("10"), True)])
        self.assertEqual(self._codes("1128K  qty 25\n1129K: 12 pieces\nT-55 - 2 nos"),
                         [("1128K", Decimal("25"), True), ("1129K", Decimal("12"), True),
                          ("T-55", Decimal("2"), True)])
        self.assertEqual(self._codes("2 nos office chair 1128K"),
                         [("1128K", Decimal("2"), True)])
        self.assertEqual(self._codes("T55 x 3 and t-55 again"),
                         [("T-55", Decimal("3"), True)])

    def test_a_code_with_no_quantity_is_not_confident(self):
        self.assertEqual(self._codes("please quote 1128K"),
                         [("1128K", None, False)])

    def test_a_price_beside_the_code_is_not_a_quantity(self):
        self.assertEqual(self._codes("price of 1128K is Rs 250 each?"),
                         [("1128K", None, False)])

    def test_words_alone_find_no_code(self):
        self.assertEqual(self._codes("chairs for the new office please"), [])

    def test_a_code_inside_a_longer_token_does_not_match(self):
        self.assertEqual(self._codes("ref 91128K5 and A1129KB"), [])


class ARateListFromAnyFile(unittest.TestCase):

    def test_word(self):
        import docx
        d = tempfile.mkdtemp()
        doc = docx.Document()
        doc.add_paragraph("ACME FURNITURE — price list 2026")
        table = doc.add_table(rows=3, cols=4)
        for r, cells in enumerate([("Code", "Description", "Unit", "Rate"),
                                   ("1128K", "Office chair mesh back", "nos", "250"),
                                   ("T-55", "Meeting table 6 seat", "nos", "9000")]):
            for c, text in enumerate(cells):
                table.cell(r, c).text = text
        p = os.path.join(d, "rates.docx")
        doc.save(p)
        items = quoting.load_rates(p)
        self.assertEqual([i.code for i in items], ["1128K", "T-55"])
        self.assertEqual(items[1].rate, Decimal("9000"))

    def test_plain_text_with_spaced_columns(self):
        d = tempfile.mkdtemp()
        p = os.path.join(d, "rates.txt")
        with open(p, "w", encoding="utf-8") as f:
            f.write("Acme price list\n\nCode    Description             Unit   Rate\n"
                    "1128K   Office chair mesh back  nos    250\n"
                    "1129K   Office chair leather    nos    400\n")
        items = quoting.load_rates(p)
        self.assertEqual([(i.code, i.rate) for i in items],
                         [("1128K", Decimal("250")), ("1129K", Decimal("400"))])

    def test_a_scanned_pdf_says_so(self):
        d = tempfile.mkdtemp()
        p = os.path.join(d, "rates.pdf")
        with open(p, "wb") as f:
            f.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        with self.assertRaises(quoting.RateFileError) as ctx:
            quoting.load_rates(p)
        self.assertIn("scan", str(ctx.exception))


class TheSetupSwitch(unittest.TestCase):

    def test_off_by_default_and_saved_when_ticked(self):
        folder = tempfile.mkdtemp()
        cfg = ready_cfg(folder)
        with _NoSave() as saved:
            d = InquirySetupDialog(cfg)
            self.assertFalse(d.auto_quote.isChecked())
            d.auto_quote.setChecked(True)
            d._save()
        self.assertTrue(saved.saved[-1]["inquiry"]["auto_quote"])
        self.assertTrue(autoquote.enabled(saved.saved[-1]))
        self.assertFalse(autoquote.enabled(ready_cfg(folder)))


class _FakeSendWorker(QObject):
    progress = Signal(int, int, str, bool, str)
    done = Signal(list, list)
    failed = Signal(str)
    sent: list = []

    def __init__(self, cfg, recipients, subject, body, files, **_):
        super().__init__()
        self.recipients, self.subject, self.body, self.files = recipients, subject, body, files
        _FakeSendWorker.sent.append(self)

    def start(self):
        self.done.emit([r["email"] for r in self.recipients], [])

    def isRunning(self):
        return False


class TheAutomaticSend(unittest.TestCase):

    def setUp(self):
        _FakeSendWorker.sent = []
        self.folder = tempfile.mkdtemp()
        self.cfg = ready_cfg(self.folder)
        self.cfg["inquiry"]["rate_list"] = _rates(self.folder)
        self.cfg["inquiry"]["auto_quote"] = True
        self.cfg["email"] = {"address": "sales@acme.co.in", "password": "x",
                             "host": "smtp.acme.co.in", "port": 465}

    def _check_with(self, body: str, *, switch=True):
        self.cfg["inquiry"]["auto_quote"] = switch
        msg = message(body=body, subject="Order enquiry")
        row = register.from_message(msg)
        paths = mailflow.Paths(self.folder)
        register.save([row], paths.register_csv)
        result = mailflow.Result()
        result.fetched = 1
        result.new_inquiries = [mailflow.Item("inquiry", message=msg, row=row)]
        with _NoSave(), \
                mock.patch.object(UI, "SendWorker", _FakeSendWorker), \
                mock.patch.object(CB.config, "save_artifact"):
            dialog = UI.InquiryDialog(self.cfg)
            dialog._refresh_register()
            dialog._checked(result, remember=False)
        return dialog, register.load(paths.register_csv)[0]

    def test_a_confident_mail_is_quoted_and_sent(self):
        dialog, row = self._check_with("Please send 40 pcs of 1128K and 1129K x 10. Thanks")
        self.assertEqual(len(_FakeSendWorker.sent), 1)
        w = _FakeSendWorker.sent[0]
        self.assertEqual(w.recipients[0]["email"], "purchase@shaktiauto.in")
        self.assertIn("Quotation", w.subject)
        self.assertIn("Office chair mesh back", w.body)
        self.assertIn("14,000.00", w.body)                 # 40×250 + 10×400
        self.assertTrue(os.path.exists(w.files[0]["path"]))
        self.assertEqual(row["Status"], register.QUOTED)
        self.assertTrue(row["Quotation no"])
        self.assertIn("quoted automatically", row["Notes"])

    def test_a_code_without_a_quantity_is_held_with_the_reason(self):
        dialog, row = self._check_with("Please quote 1128K and 1129K x 10")
        self.assertEqual(_FakeSendWorker.sent, [])
        self.assertEqual(row["Status"], register.NEW)
        self.assertIn("auto-quote held", row["Notes"])
        self.assertIn("1128K", row["Notes"])

    def test_a_mail_in_words_only_is_held(self):
        dialog, row = self._check_with("Kindly quote 5000 nos compression spring 2mm.")
        self.assertEqual(_FakeSendWorker.sent, [])
        self.assertIn("no product code", row["Notes"])

    def test_the_switch_off_sends_nothing(self):
        dialog, row = self._check_with("40 pcs of 1128K please", switch=False)
        self.assertEqual(_FakeSendWorker.sent, [])
        self.assertEqual(row["Status"], register.NEW)
        self.assertNotIn("auto-quote", row["Notes"])


class TheQuotationWindowShowsTheLines(unittest.TestCase):

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.cfg = ready_cfg(self.folder)
        self.cfg["inquiry"]["rate_list"] = _rates(self.folder)
        self.items = quoting.load_rates(self.cfg["inquiry"]["rate_list"])

    def _window(self, body: str):
        row = register.from_message(message(body=body))
        row["Product asked"] = body
        register.save([row], mailflow.Paths(self.folder).register_csv)
        with _NoSave():
            dialog = UI.InquiryDialog(self.cfg)
            dialog._refresh_register()
            return UI.QuotationDialog(self.cfg, dialog._register_rows[0],
                                      self.items, dialog)

    def test_two_codes_make_a_lines_table_and_a_two_line_quote(self):
        q = self._window("1128K x 40 and 1129K x 10")
        self.assertTrue(q.lines_table.isVisibleTo(q))
        self.assertFalse(q.single_form.isVisibleTo(q))
        self.assertEqual(q.lines_table.rowCount(), 2)
        self.assertEqual(len(q.quote.lines), 2)
        self.assertEqual(q.quote.subtotal, Decimal("14000.00"))
        self.assertIn("Office chair leather", q.preview.toPlainText())
        q.lines_table.item(0, 2).setText("50")
        self.assertEqual(q.quote.subtotal, Decimal("16500.00"))

    def test_one_code_prefills_the_form(self):
        q = self._window("please send 1129K x 10")
        self.assertFalse(q.lines_table.isVisibleTo(q))
        self.assertEqual(q.item_picker.currentData().code, "1129K")
        self.assertEqual(q.quantity.text(), "10")
        self.assertEqual(q.quote.subtotal, Decimal("4000.00"))


if __name__ == "__main__":
    unittest.main()
