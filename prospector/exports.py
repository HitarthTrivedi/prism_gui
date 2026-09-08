"""
Prism Sales Automation — the deliverables
──────────────────────────────────────────
The two spreadsheets a seller actually hands over, at the quality of a
hand-built list rather than a flat dump:

  · LEADS — every row, industry-tabbed, its e-mail verified against a live MX
    lookup and colour-coded so "which of these can I actually reach" is one
    glance, not a guess.
  · HOT LIST — the qualified case files: verdict, score, the six named
    dimensions (each with its evidence and source) and the drafted why-now
    opener, colour-coded by verdict.

openpyxl only — the same library the BOQ pricing add-on already ships — and one
MX lookup per UNIQUE domain (a real list repeats a company many times, and the
answer is the same each time), so a 2,000-row sheet is a few hundred lookups.
"""
from __future__ import annotations

import os
import re
from collections import OrderedDict, defaultdict
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import COLD, HOT, WARM, Dossier

# Colour-coding shared by both sheets, lifted from the hand-built VITRACX files.
_EFILL = {"valid": "E4F2EA", "domain-ok": "E4F2EA",
          "invalid": "FBD5D5", "no-mx": "FBE6E4",
          "catch-all": "FBEFD5", "unknown": "F3F1E4", "no-domain": "F0F0F0"}
_VFILL = {"HOT": "FBDAD5", "WARM": "FBEFD5", "COLD": "ECECEC"}
_HEAD_BLUE = "243B53"


# ── e-mail verification (MX only — cheap, no paid vendor) ─────────────────────

def _domain(email: str) -> str:
    return email.split("@", 1)[1].lower().strip() if email and "@" in email else ""


def build_email_checks(leads) -> dict:
    """{domain: 'domain-ok'|'no-mx'|'unknown'} for every domain in the list, one
    threaded DNS-over-HTTPS lookup per UNIQUE domain."""
    import requests  # lazy, like the rest of the engine
    uniq = sorted({_domain(getattr(l, "email", "")) for l in leads
                   if _domain(getattr(l, "email", ""))})
    if not uniq:
        return {}

    def one(dom: str) -> str:
        try:
            j = requests.get("https://dns.google/resolve",
                             params={"name": dom, "type": "MX"}, timeout=5).json()
            return "domain-ok" if j.get("Answer") else "no-mx"
        except Exception:                               # noqa: BLE001
            return "unknown"

    with ThreadPoolExecutor(max_workers=32) as ex:
        return dict(zip(uniq, ex.map(one, uniq)))


def check_of(lead, mx: dict) -> str:
    # A verifier's status on the lead (Hunter) wins over the domain-only MX ping.
    ec = (getattr(lead, "extra", None) or {}).get("email_check")
    if ec:
        return ec
    e = getattr(lead, "email", "")
    if not e:
        return "no-domain"
    return mx.get(_domain(e), "unknown")


# ── shared styling helpers ────────────────────────────────────────────────────

def _fill(hexrgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=hexrgb)


def _safe_title(name: str, used: set) -> str:
    """A worksheet title openpyxl will accept: <=31 chars, none of []:*?/\\, and
    unique within the workbook."""
    t = re.sub(r"[\[\]:*?/\\]", " ", (name or "Leads")).strip()[:31] or "Leads"
    base, n = t, 2
    while t.lower() in used:
        t = f"{base[:28]} {n}"
        n += 1
    used.add(t.lower())
    return t


def _save(wb, path: str) -> str:
    """Save, but never let an ALREADY-OPEN file break a run. Windows locks a
    spreadsheet the user has open in Excel/WPS, so a re-run's overwrite throws
    PermissionError. Fall back to a timestamped sibling and return the path
    actually written, so the run still delivers a sheet instead of erroring."""
    try:
        wb.save(path)
        return path
    except (PermissionError, OSError):
        import datetime
        base, ext = os.path.splitext(path)
        alt = f"{base} {datetime.datetime.now():%H%M%S}{ext}"
        wb.save(alt)
        return alt


# ── the LEADS sheet — industry-tabbed, e-mail-verified ────────────────────────

_LEAD_COLS = ["No.", "Industry", "Company", "Name", "Designation", "E-mail",
              "Email check", "Phone No.", "Location", "In role since", "LinkedIn"]
_LEAD_WIDTH = [5, 20, 30, 22, 44, 34, 12, 13, 22, 12, 34]


def leads_xlsx(leads, path: str) -> str:
    """Every lead, grouped into one tab per industry, with a live MX check on
    each e-mail. The block columns (No./Industry/Company) are shown once per
    company and forward-filled visually, exactly like the hand-built sheet."""
    mx = build_email_checks(leads)
    head_fill, head_font = _fill("DCE3EF"), Font(bold=True, color=_HEAD_BLUE)
    bold, top = Font(bold=True), Border(top=Side(style="thin", color="C7CDD6"))
    al = Alignment(vertical="top")

    by_ind: "OrderedDict[str, defaultdict]" = OrderedDict()
    for l in leads:
        ind = (l.industry or "Leads").strip() or "Leads"
        by_ind.setdefault(ind, defaultdict(list))[(l.company or "").strip()].append(l)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set = set()
    for ind, comps in by_ind.items():
        ws = wb.create_sheet(title=_safe_title(ind, used))
        for ci, name in enumerate(_LEAD_COLS, 1):
            c = ws.cell(1, ci, name)
            c.fill, c.font, c.alignment = head_fill, head_font, al
            ws.column_dimensions[get_column_letter(ci)].width = _LEAD_WIDTH[ci - 1]
        ws.freeze_panes = "A2"
        rn, cno = 2, 0
        for company, ppl in sorted(comps.items(), key=lambda kv: kv[0].lower()):
            cno += 1
            for pi, l in enumerate(ppl):
                first = pi == 0
                x = l.extra or {}
                vals = [cno if first else "", ind if first else "",
                        company if first else "", l.name, l.title, l.email,
                        check_of(l, mx), l.phone, x.get("location", ""),
                        x.get("since", ""), x.get("linkedin", "")]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(rn, ci, v)
                    c.alignment = al
                    if first:
                        c.border = top
                        if ci <= 3:
                            c.font = bold
                    if _LEAD_COLS[ci - 1] == "Email check" and v in _EFILL:
                        c.fill = _fill(_EFILL[v])
                rn += 1
    return _save(wb, path)


# ── the HOT LIST — the qualified case files ───────────────────────────────────

_DIMS = ["Fit", "Need", "Timing", "Authority", "Budget", "Competition"]
_HOT_COLS = (["Verdict", "Score", "Company", "Name", "Designation", "E-mail",
              "Email check", "Industry"] + _DIMS + ["Why-now opener"])
_HOT_WIDTH = [8, 6, 24, 20, 32, 30, 12, 14] + [32] * 6 + [54]
_RANK = {HOT: 0, WARM: 1, COLD: 2}


def hotlist_xlsx(dossiers, path: str) -> str:
    """The qualified dossiers as one colour-coded sheet: verdict, score, the six
    dimensions (mark · evidence · [source]) and the drafted opener."""
    rows = sorted(dossiers, key=lambda d: (_RANK.get(d.verdict, 3), -d.score))
    mx = build_email_checks([d.lead for d in rows])
    head_fill, head_font = _fill(_HEAD_BLUE), Font(bold=True, color="FFFFFF")
    wrap = Alignment(vertical="top", wrap_text=True)
    topl = Alignment(vertical="top")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hot list"
    for ci, name in enumerate(_HOT_COLS, 1):
        c = ws.cell(1, ci, name)
        c.fill, c.font, c.alignment = head_fill, head_font, topl
        ws.column_dimensions[get_column_letter(ci)].width = _HOT_WIDTH[ci - 1]
    ws.freeze_panes = "C2"

    for ri, d in enumerate(rows, 2):
        dims = {dim.name: dim for dim in d.dimensions}
        verdict = d.verdict.upper()
        cells = [verdict, d.score, d.lead.company, d.lead.name, d.lead.title,
                 d.lead.email, check_of(d.lead, mx), d.lead.industry]
        for name in _DIMS:
            dim = dims.get(name)
            if dim and (dim.evidence or dim.source):
                src = f"  [{dim.source}]" if dim.source else ""
                cells.append(f"{dim.mark()} {dim.evidence}{src}".strip())
            else:
                cells.append(dim.mark() if dim else "?")
        cells.append(d.opener)
        for ci, v in enumerate(cells, 1):
            c = ws.cell(ri, ci, v)
            c.alignment = wrap if _HOT_COLS[ci - 1] in _DIMS + ["Why-now opener"] else topl
            if ci == 1 and verdict in _VFILL:
                c.fill, c.font = _fill(_VFILL[verdict]), Font(bold=True)
            if _HOT_COLS[ci - 1] == "Email check" and v in _EFILL:
                c.fill = _fill(_EFILL[v])
        ws.row_dimensions[ri].height = 58
    return _save(wb, path)
